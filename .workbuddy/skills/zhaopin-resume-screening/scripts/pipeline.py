"""
zhaopin-resume-screening — 评估管线入口（一条命令跑完机械部分）。

把"复制模版 → 提取简历 → 按模版填表 → 写入 AI 评估 → 读回校验"的固定流程
固化为代码；Agent 只需产出评估 JSON（B 评级 + C 评价）。

用法（在项目 venv 下运行）:
    # 阶段1: 准备评估表（复制模版 + 提取 + 填基础列），并生成待评估清单
    python pipeline.py prepare --resume-dir <docx目录> --output-xlsx <输出.xlsx> [--workdir <中间产物目录>]

    # 阶段2: AI 评估完成后写入 B/C 列（evaluations.json 见下）并读回校验
    python pipeline.py evaluate --xlsx <输出.xlsx> --evaluations <evaluations.json> [--workdir <中间产物目录>]

evaluations.json 格式（index 为数据行号，0 计；Excel 行 = index + 2）:
    [
      {"index": 0, "rating": "高", "comment": "12年销售经验... 薪资期望 8-12k vs 岗位 6-10k"},
      {"index": 1, "rating": "低", "comment": "..."}
    ]
rating 取值: 高 / 中 / 低（对应色阶 C6EFCE / FFEB9C / FFC7CE）

中间产物（workdir，默认与 output-xlsx 同目录）:
    extracted_data.json  提取结果（原始顺序）
    rows.json            填表后 Excel 行顺序（index → filename 映射，评估依据）
"""

import argparse
import json
import shutil
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
TEMPLATE_PATH = SKILL_DIR / "templates" / "EvaluationTemplate.xlsx"

sys.path.insert(0, str(SCRIPT_DIR))

from extract_docx import batch_extract          # noqa: E402
from fill_excel import fill_excel_template      # noqa: E402
from openpyxl import load_workbook              # noqa: E402

VALID_RATINGS = {"高", "中", "低"}
RATING_COLORS = {"高": "C6EFCE", "中": "FFEB9C", "低": "FFC7CE"}


def cmd_prepare(resume_dir: str, output_xlsx: str, workdir: str | None) -> int:
    resume_dir = Path(resume_dir).resolve()
    output_xlsx = Path(output_xlsx).resolve()
    workdir = Path(workdir).resolve() if workdir else output_xlsx.parent
    workdir.mkdir(parents=True, exist_ok=True)

    if not resume_dir.is_dir():
        print(f"[FAIL] resume dir not found: {resume_dir}")
        return 1
    if not TEMPLATE_PATH.exists():
        print(f"[FAIL] template not found: {TEMPLATE_PATH}")
        return 1

    docx_files = sorted(resume_dir.glob("*.docx"))
    if not docx_files:
        print(f"[FAIL] no .docx files in {resume_dir}")
        return 1

    # 1) 复制模版（fill_excel_template 是就地覆盖，必须先复制）
    if output_xlsx.exists() and output_xlsx.resolve() == TEMPLATE_PATH.resolve():
        print("[FAIL] output path equals template path - refuse to overwrite")
        return 1
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE_PATH, output_xlsx)
    print(f"[OK] template copied -> {output_xlsx}")

    # 2) 提取
    data = batch_extract(str(resume_dir))
    extracted_json = workdir / "extracted_data.json"
    with open(extracted_json, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"[OK] extracted {len(data)} resumes -> {extracted_json}")

    # 3) 按模版填充（内置 姓名+年龄 去重；B/C 留空）
    filled = fill_excel_template(str(extracted_json), str(output_xlsx))
    print(f"[OK] filled {filled} rows (dedup skipped {len(data) - filled})")

    # 4) 导出最终行顺序 —— AI 评估必须以此 index 为准（勿按姓名匹配，同名同龄会串行）
    wb = load_workbook(str(output_xlsx))
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
    rows = []
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        if v:
            rows.append({"index": row - 2, "excel_row": row, "filename": str(v)})
    rows_json = workdir / "rows.json"
    with open(rows_json, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    print(f"[OK] row order exported ({len(rows)} rows) -> {rows_json}")
    print(f"[NEXT] let the AI evaluate each rows.json entry, write evaluations.json, then run:")
    print(f"       python pipeline.py evaluate --xlsx \"{output_xlsx}\" --evaluations <evaluations.json>")
    return 0


def cmd_evaluate(xlsx: str, evaluations: str, workdir: str | None) -> int:
    from fill_excel import write_evaluation_row

    xlsx_path = Path(xlsx).resolve()
    workdir = Path(workdir).resolve() if workdir else xlsx_path.parent
    if not xlsx_path.exists():
        print(f"[FAIL] xlsx not found: {xlsx_path}")
        return 1
    with open(evaluations, "r", encoding="utf-8") as f:
        evals = json.load(f)

    # 以 rows.json 为行序基准（若无则按当前表格顺序）
    rows_json = workdir / "rows.json"
    if rows_json.exists():
        with open(rows_json, "r", encoding="utf-8") as f:
            rows = json.load(f)
    else:
        print("[WARN] rows.json not found - deriving row order from xlsx column A")
        wb0 = load_workbook(str(xlsx_path))
        ws0 = wb0["Sheet1"] if "Sheet1" in wb0.sheetnames else wb0.active
        rows = []
        for row in range(2, ws0.max_row + 1):
            v = ws0.cell(row=row, column=1).value
            if v:
                rows.append({"index": row - 2, "excel_row": row, "filename": str(v)})

    written = 0
    invalid = []
    for ev in evals:
        idx = ev.get("index")
        rating = str(ev.get("rating", "")).strip()
        comment = str(ev.get("comment", ""))
        if idx is None or not isinstance(idx, int) or idx < 0 or idx >= len(rows):
            invalid.append({"reason": "bad index", "entry": ev})
            continue
        if rating not in VALID_RATINGS:
            invalid.append({"reason": f"rating must be one of {sorted(VALID_RATINGS)}", "entry": ev})
            continue
        write_evaluation_row(str(xlsx_path), idx, rating, comment)
        written += 1
    print(f"[OK] wrote {written} evaluation rows -> {xlsx_path}")
    for bad in invalid:
        print(f"[SKIP] {bad['reason']}: {json.dumps(bad['entry'], ensure_ascii=False)[:120]}")

    return cmd_verify(str(xlsx_path))


def cmd_verify(xlsx: str) -> int:
    xlsx_path = Path(xlsx).resolve()
    wb = load_workbook(str(xlsx_path))
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

    total = 0
    rated = {"高": 0, "中": 0, "低": 0}
    empty_bc = []
    samples = []
    for row in range(2, ws.max_row + 1):
        a = ws.cell(row=row, column=1).value
        if not a:
            continue
        total += 1
        b = ws.cell(row=row, column=2).value
        c = ws.cell(row=row, column=3).value
        if b in rated:
            rated[b] += 1
        else:
            empty_bc.append(row)
        if len(samples) < 3:
            samples.append({
                "excel_row": row,
                "filename": str(a)[:50],
                "rating": b,
                "comment": (str(c)[:60] + "…") if c and len(str(c)) > 60 else c,
            })

    print("=== VERIFY ===")
    print(f"data rows          : {total}")
    print(f"rating distribution: 高={rated['高']} 中={rated['中']} 低={rated['低']}")
    print(f"rows missing B/C   : {len(empty_bc)}"
          + (f" -> excel rows {empty_bc[:20]}" if empty_bc else " (all rated)"))
    print(f"samples            : {json.dumps(samples, ensure_ascii=False)}")
    ok = len(empty_bc) == 0 and total > 0
    print(f"[{'OK' if ok else 'WARN'}] verify {'passed' if ok else 'found gaps - check above'}")
    return 0 if ok else 2


def main() -> int:
    ap = argparse.ArgumentParser(description="zhaopin resume screening pipeline")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p1 = sub.add_parser("prepare", help="copy template + extract + fill base columns")
    p1.add_argument("--resume-dir", required=True)
    p1.add_argument("--output-xlsx", required=True)
    p1.add_argument("--workdir", default=None)

    p2 = sub.add_parser("evaluate", help="write AI evaluations (B/C) + verify")
    p2.add_argument("--xlsx", required=True)
    p2.add_argument("--evaluations", required=True)
    p2.add_argument("--workdir", default=None)

    p3 = sub.add_parser("verify", help="read back and verify xlsx")
    p3.add_argument("--xlsx", required=True)

    args = ap.parse_args()
    if args.cmd == "prepare":
        return cmd_prepare(args.resume_dir, args.output_xlsx, args.workdir)
    if args.cmd == "evaluate":
        return cmd_evaluate(args.xlsx, args.evaluations, args.workdir)
    if args.cmd == "verify":
        return cmd_verify(args.xlsx)
    return 1


if __name__ == "__main__":
    sys.exit(main())
