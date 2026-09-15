"""
从智联招聘 DOCX 简历中提取关键信息。
- 段落含: 基本信息、求职意向、教育经历
- 表格含: 工作经历（Table 1, 单列格式）
每组: Company+Position → Salary+Time → Description → Empty

Usage:
    from extract_docx import batch_extract
    results = batch_extract("/path/to/docx/folder")
"""

import re
import json
from datetime import datetime
from pathlib import Path
from docx import Document


def extract_text_from_docx(docx_path):
    """返回 (段落文本, 表格列表)"""
    doc = Document(str(docx_path))
    para_text = "\n".join([p.text for p in doc.paragraphs])
    return para_text, doc.tables


def parse_name(para_text):
    """从段落中提取姓名。格式: 'XXX点击此处联系候选人'"""
    for line in para_text.strip().split('\n'):
        ls = line.strip()
        if not ls:
            continue
        m = re.match(r'^(.+?)点击此处', ls)
        if m:
            return m.group(1).strip()
    return "未知"


def parse_basic_info(para_text):
    """从段落提取性别、年龄、现居地。"""
    result = {}

    m = re.search(r'[｜|]\s*(男|女)\s*[｜|]', para_text)
    if m:
        result['gender'] = m.group(1)
    else:
        m = re.search(r'(?<![a-zA-Z])(男|女)(?![a-zA-Z])', para_text)
        result['gender'] = m.group(1) if m else ""

    m = re.search(r'(\d+)\s*岁', para_text)
    age = int(m.group(1)) if m else 0
    if age == 0:
        m = re.search(r'\((\d{4})\s*年', para_text)
        if m:
            age = datetime.now().year - int(m.group(1))
    result['age'] = age

    m = re.search(r'现居住地[：:]\s*(.+?)(?:[｜|\n])', para_text)
    result['current_residence'] = m.group(1).strip() if m else ""

    return result


def parse_salary_range(salary_str):
    """Parse salary string to (low, high) in 千元/月.
    Examples: '1.5万-2万/月' → (15, 20), '8千-1.5万/月' → (8, 15), '15k-25k' → (15, 25)
    Returns (None, None) if unparseable."""
    if not salary_str:
        return None, None
    salary_str = salary_str.replace(' ', '')
    # Normalize: 万 → *10, 千 → *1, k/K → *1
    def to_k(val_str):
        val_str = val_str.replace(',', '')
        return float(val_str)

    # Match range: X万-Y万/月, X千-Y千/月, Xk-Yk, etc.
    m = re.search(r'([\d,.]+)\s*([万千kK]?)\s*-\s*([\d,.]+)\s*([万千kK]?)(?:/(?:月|年))?', salary_str)
    if m:
        v1, u1, v2, u2 = m.group(1), m.group(2).lower(), m.group(3), m.group(4).lower()
        n1 = to_k(v1)
        n2 = to_k(v2)
        if u1 in ('万', 'w') or (not u1 and u2 in ('万', 'w')):
            n1 *= 10
        if u2 in ('万', 'w') or (not u2 and u1 in ('万', 'w')):
            n2 *= 10
        if '年' in salary_str:
            n1 = round(n1 / 12)
            n2 = round(n2 / 12)
        return int(n1), int(n2)

    # Match single value: X万/月, X千/月
    m = re.search(r'([\d,.]+)\s*([万千kK]?)(?:/(?:月|年))?', salary_str)
    if m:
        v, u = m.group(1), m.group(2).lower()
        n = to_k(v)
        if u in ('万', 'w'):
            n *= 10
        if '年' in salary_str:
            n = round(n / 12)
        return int(n), int(n)

    return None, None


def parse_job_intentions(para_text):
    """从段落提取求职意向。返回 (display_text, salary_info_list)
    salary_info_list: [{position, salary_low_k, salary_high_k, location, industry}, ...]"""
    intentions = []
    salary_info_list = []

    m = re.search(r'求职意向\s*\n(.*?)(?=\n\s*(?:工作经历|项目经历|$))', para_text, re.DOTALL)
    if not m:
        return "", []
    intention_raw = m.group(1)

    parts = re.split(r'(全职|兼职(?:/临时)?)', intention_raw)
    intent_pairs, current = [], ""
    for part in parts:
        if part in ('全职', '兼职', '兼职/临时'):
            intent_pairs.append((current.strip(), part))
            current = ""
        else:
            current = part
    if current.strip():
        intent_pairs.append((current.strip(), ""))

    for content, _job_type in intent_pairs:
        content = content.strip()
        if not content:
            continue
        lines = [l.strip() for l in content.split('\n') if l.strip()]
        if not lines:
            continue

        first = lines[0]
        # Split position and location by 2+ spaces
        parts_pos = re.split(r'\s{2,}', first, maxsplit=1)
        position = parts_pos[0].replace(' ', '')
        location = parts_pos[1].strip() if len(parts_pos) > 1 else ""

        salary = ""
        industry = ""

        for sl in lines[1:]:
            sl_clean = sl.strip().replace(' ', '')
            if not sl_clean:
                continue
            sal_m = re.search(r'([\d,.]+[万千百元](?:-[\d,.]*[万千百元])?\s*/\s*[年月])', sl_clean)
            if not sal_m:
                sal_m = re.search(r'([\d,.]+[万千百元])', sl_clean)
            if sal_m:
                salary = re.sub(r'\s+', '', sal_m.group(1))
                continue
            if '行业' in sl_clean:
                industry = sl_clean
            elif '不限' in sl_clean:
                industry = '不限行业'
            elif not location and re.match(r'^[\u4e00-\u9fa5]{2,8}$', sl_clean):
                location = sl_clean

        if not industry:
            industry = "不限行业"

        if position and len(position) >= 2 and not re.match(r'^[、，,./]+$', position):
            intentions.append(f"【{position}】{salary}-{location}-{industry}")
            low_k, high_k = parse_salary_range(salary)
            salary_info_list.append({
                'position': position,
                'salary_low_k': low_k,
                'salary_high_k': high_k,
                'location': location,
                'industry': industry,
            })

    return '\n'.join(intentions) if intentions else "", salary_info_list


def parse_work_experiences(tables):
    """从表格提取工作经历。每组: Company+Position → Salary+Time → Desc → Empty
    返回 (display_text, descriptions_text)：
    - display_text: 写入 Excel H 列，仅含【岗位】公司-时间（不含工作描述）
    - descriptions_text: 所有工作描述的合并文本，供 AI 行业匹配度判断"""
    experiences = []
    all_descriptions = []

    for table in tables:
        if len(table.columns) != 1:
            continue

        rows = [row.cells[0].text.strip() for row in table.rows]

        # Skip metadata tables (应聘职位, 期望从事, 驾驶证, 技能等)
        if rows and any(kw in rows[0] for kw in ['应聘', '期望从事', '驾驶', '硬件']):
            continue

        i = 0
        while i < len(rows):
            line = rows[i]
            if not line or '工作描述' in line:
                i += 1
                continue

            # Match: "{Company}   {Position}  ({Alias})"
            co_pos_match = re.match(r'^\s*(.+?)\s{2,}(.+?)(?:\s*\(([^)]*)\))?\s*$', line)
            if not co_pos_match:
                i += 1
                continue

            company = co_pos_match.group(1).strip()
            position = co_pos_match.group(2).strip().replace(' ', '')

            if not company or not re.search(r'[\u4e00-\u9fa5a-zA-Z]', company):
                i += 1
                continue

            # Verify company line (not salary line): must not start with digits
            if re.match(r'^[\d,.]+[万千百元]', company):
                i += 1
                continue

            # Next row = salary + time
            time_range = ""
            if i + 1 < len(rows):
                next_line = rows[i + 1]
                tm = re.search(r'(\d{4}\.\d{1,2}\s*-\s*(?:至今|\d{4}\.\d{1,2}))', next_line)
                if tm:
                    time_range = tm.group(1).replace(' ', '')
                dur_m = re.search(r'\((\d+年\d*个?月?|\d+个月|\d+年)\)', next_line)
                if dur_m and time_range:
                    time_range += dur_m.group(0).replace(' ', '')

            # Next+2 row = work description (only for AI evaluation, not written to Excel)
            if i + 2 < len(rows):
                desc_line = rows[i + 2]
                if desc_line.startswith('工作描述'):
                    desc_content = re.sub(r'^工作描述[：:]\s*', '', desc_line)
                    if desc_content:
                        all_descriptions.append(f"【{position}】{company}: {desc_content[:300]}")

            # Excel H column: position + company + time only, no description
            formatted = f"【{position}】{company}"
            if time_range:
                formatted += f"-{time_range}"
            experiences.append(formatted)

            # Skip: salary+time row, description row, empty separator (3 rows)
            i += 4 if i + 3 < len(rows) else 1

    # Deduplicate
    seen = set()
    unique = []
    for exp in experiences:
        key = exp[:40]
        if key not in seen:
            seen.add(key)
            unique.append(exp)

    display = '\n'.join(unique) if unique else ""
    desc = '\n'.join(all_descriptions) if all_descriptions else ""
    return display, desc


def parse_resume(docx_path):
    """解析单份 DOCX 简历，返回结构化字典。
    work_experiences: 写入 Excel H 列（公司+岗位+时间）
    work_descriptions: 工作描述文本（仅用于 AI 行业匹配度判断，不写入 Excel）
    salary_info: 结构化薪资信息列表（仅用于 AI 薪资对比，不写入 Excel）"""
    para_text, tables = extract_text_from_docx(docx_path)

    result = {'name': parse_name(para_text)}
    result.update(parse_basic_info(para_text))
    job_intentions_text, salary_info = parse_job_intentions(para_text)
    result['job_intentions'] = job_intentions_text
    result['salary_info'] = salary_info
    display, desc = parse_work_experiences(tables)
    result['work_experiences'] = display
    result['work_descriptions'] = desc

    return result


def batch_extract(docx_dir):
    """批量提取目录下所有 DOCX 简历。"""
    docx_files = sorted(Path(docx_dir).glob("*.docx"))
    if not docx_files:
        raise FileNotFoundError(f"No DOCX files found in {docx_dir}")

    results = []
    for docx_file in docx_files:
        parsed = parse_resume(docx_file)
        parsed['filename'] = docx_file.name
        results.append(parsed)

    return results


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python extract_docx.py <docx_directory> [output_json]")
        sys.exit(1)

    docx_dir = sys.argv[1]
    output = sys.argv[2] if len(sys.argv) > 2 else "extracted_data.json"

    results = batch_extract(docx_dir)
    with open(output, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"Extracted {len(results)} resumes -> {output}")
